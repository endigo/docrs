"""Oracle extension for work package p5 (M5 review workflows).

RESOLVERS: {"$tag": fn(document, value)} argument decoders.
ENCODERS: [fn(value) -> (handled, encoded)] result encoders.
FUNCTIONS: {"name": fn(document, *args, **kwargs)} callable as op "ext.name".

Method-style upstream APIs (`document.revisions.accept_all()`,
`Revision.accept()`, `document.add_comment(runs, ...)`, `Span.comment(...)`,
`Comment.add_paragraph(...)`) are exposed as `ext.p5_*` helpers; the Rust
adapters in `crates/docrs-compat/src/ops/{revision,commentops}.rs`
register the same names.
"""


def _select(revisions, select):
    """A revision by index, or by filters {type, text, mark, author, nth}."""
    if select is None:
        select = 0
    if isinstance(select, int):
        return revisions[select]
    items = list(revisions)
    if "type" in select:
        items = [r for r in items if r.revision_type == select["type"]]
    if "text" in select:
        items = [r for r in items if select["text"] in r.text]
    if "mark" in select:
        items = [r for r in items if r.is_paragraph_mark == select["mark"]]
    if "author" in select:
        items = [r for r in items if r.author == select["author"]]
    return items[select.get("nth", 0)]


def p5_revisions(document):
    return document.revisions.to_dict()


def p5_remaining_unsupported(document):
    return document.revisions.remaining_unsupported()


def p5_accept_all(document, author=None):
    return document.revisions.accept_all(author=author)


def p5_reject_all(document, author=None):
    return document.revisions.reject_all(author=author)


def p5_accept(document, select=0):
    _select(document.revisions, select).accept()


def p5_reject(document, select=0):
    _select(document.revisions, select).reject()


def p5_remaining_markup(document):
    from docx.revision import _remaining_markup

    return _remaining_markup(document)


def p5_stale(document, first=0, action="accept_all", index=1):
    """Resolve revision `first` through one snapshot, then retry `action` on
    that (now stale) snapshot."""
    snapshot = document.revisions
    items = list(snapshot)
    items[first].accept()
    if action == "accept_all":
        return snapshot.accept_all()
    if action == "reject_all":
        return snapshot.reject_all()
    if action == "accept":
        return items[index].accept()
    if action == "reject":
        return items[index].reject()
    raise ValueError(action)


def p5_accept_all_twice(document):
    """`accept_all` refreshes its own snapshot: a second call resolves 0."""
    snapshot = document.revisions
    first = snapshot.accept_all()
    return [first, snapshot.accept_all(), len(snapshot)]


def p5_tracked_replace(document, span, text, author, date=None):
    span.replace(text, tracked=True, author=author, date=date)


def p5_replace(document, span, text):
    span.replace(text)


# -- comments ---------------------------------------------------------------


def _comments_part(document):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    return document.part.part_related_by(RT.COMMENTS)


def _comment(document, comment_id):
    for comment in _comments_part(document).comments:
        if comment.comment_id == comment_id:
            return comment
    raise LookupError(f"no comment {comment_id}")


def _comment_dict(comment):
    timestamp = comment.timestamp
    return {
        "comment_id": comment.comment_id,
        "author": comment.author,
        "initials": comment.initials,
        "timestamp": timestamp.isoformat() if timestamp is not None else None,
        "text": comment.text,
    }


def p5_comments(document):
    try:
        part = _comments_part(document)
    except KeyError:
        return []
    return [_comment_dict(comment) for comment in part.comments]


def _span_runs(span):
    from docx.text.run import Run

    runs = []
    for atom in span._atoms:  # noqa: SLF001
        if atom.run is not None and not any(atom.run is run for run in runs):
            runs.append(atom.run)
    return [Run(run, span._document._body) for run in runs]  # noqa: SLF001


def p5_add_comment(document, span, text="", author="", initials="", date=None):
    comment = document.add_comment(
        _span_runs(span), text=text, author=author, initials=initials
    )
    if date is not None:
        comment._comment_elm.date = date  # noqa: SLF001
    return _comment_dict(comment)


def p5_comments_add(document, text="", author="", initials="", date=None):
    comment = document.comments.add_comment(text, author=author, initials=initials)
    if date is not None:
        comment._comment_elm.date = date  # noqa: SLF001
    return _comment_dict(comment)


def p5_span_comment(document, span, text, author, initials=None, date=None):
    return _comment_dict(span.comment(text, author=author, initials=initials, date=date))


def p5_reply(document, comment_id, text, author, initials=None, date=None):
    from docx.commentops import reply

    return _comment_dict(
        reply(
            document,
            _comment(document, comment_id),
            text,
            author=author,
            initials=initials,
            date=date,
        )
    )


def p5_resolve(document, comment_id, resolved=True):
    from docx.commentops import resolve

    resolve(document, _comment(document, comment_id), resolved=resolved)


def p5_is_resolved(document, comment_id):
    from docx.commentops import is_resolved

    return is_resolved(document, _comment(document, comment_id))


def p5_parent_of(document, comment_id):
    from docx.commentops import parent_of

    return parent_of(document, _comment(document, comment_id))


def p5_anchored_text(document, comment_id):
    from docx.commentops import anchored_text

    return anchored_text(document, _comment(document, comment_id))


def p5_delete_comment(document, comment_id):
    from docx.commentops import delete_comment

    delete_comment(document, _comment(document, comment_id))


def p5_comment_add_paragraph(document, comment_id, text=""):
    comment = _comment(document, comment_id)
    comment.add_paragraph(text)
    return _comment_dict(comment)


def p5_comment_add_table(document, comment_id, rows, cols, width_emu):
    from docx.shared import Emu

    comment = _comment(document, comment_id)
    comment.add_table(rows, cols, Emu(width_emu))
    return _comment_dict(comment)


def p5_comment_set_author(document, comment_id, author):
    comment = _comment(document, comment_id)
    comment.author = author
    return _comment_dict(comment)


def p5_comment_set_initials(document, comment_id, initials):
    comment = _comment(document, comment_id)
    comment.initials = initials
    return _comment_dict(comment)


RESOLVERS = {}
ENCODERS = []
FUNCTIONS = {
    name: value
    for name, value in dict(globals()).items()
    if name.startswith("p5_") and callable(value)
}
