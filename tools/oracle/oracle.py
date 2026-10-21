"""Pinned Python oracle adapter (development only; never a runtime dependency).

Usage:
    .venv/bin/python -I tools/oracle/oracle.py REQUESTS.json > RESULTS.json

REQUESTS.json is a JSON array of request objects; the output is a JSON array
of results in the same order. Each request names an `op`; results are
normalized to the vocabulary the Rust harness compares:

    {"outcome": "ok", ...op-specific fields...}
    {"outcome": "error", "class": "<ExceptionClass>", "refusal": bool,
     "category": "<rust-style code or 'non_refusal_error'>"}

The oracle is the pinned upstream paper-docx 0.2.1 installed in .venv.
"""

from __future__ import annotations

import hashlib
import json
import sys
import traceback


def _category(exc: BaseException) -> str:
    from docx import errors

    mapping = [
        (errors.MalformedPackageError, "malformed_package"),
        (errors.AmbiguousTargetError, "ambiguous_target"),
        (errors.TargetNotFoundError, "target_not_found"),
        (errors.UnsupportedStructureError, "unsupported_structure"),
        (errors.BoundaryViolationError, "boundary_violation"),
        (errors.RelationshipPolicyError, "relationship_policy"),
        (errors.DocumentProtectedError, "document_protected"),
    ]
    for cls, code in mapping:
        if isinstance(exc, cls):
            return code
    if isinstance(exc, (TypeError, ValueError)) and not isinstance(exc, errors.PaperRefusal):
        return "invalid_argument"
    return "non_refusal_error"


def _error(exc: BaseException) -> dict:
    from docx.errors import PaperRefusal

    return {
        "outcome": "error",
        "class": type(exc).__name__,
        "refusal": isinstance(exc, PaperRefusal),
        "category": _category(exc),
        "message": str(exc)[:400],
    }


def op_open(request: dict) -> dict:
    import docx

    docx.Document(request["path"])
    return {"outcome": "ok"}


def op_diff(request: dict) -> dict:
    from docx.package import diff_package

    return {"outcome": "ok", "diff": diff_package(request["a"], request["b"]).to_dict()}


def op_noop_patch_save(request: dict) -> dict:
    import docx
    from docx.package import patch_save

    document = docx.Document(request["path"])
    result = patch_save(request["path"], document, request["out"])
    with open(request["out"], "rb") as handle:
        out_sha = hashlib.sha256(handle.read()).hexdigest()
    return {"outcome": "ok", "report": result.to_dict(), "out_sha256": out_sha}


def op_diagnose(request: dict) -> dict:
    from docx.package import diagnose

    return {"outcome": "ok", "diagnosis": diagnose(request["path"]).to_dict()}


def _span_dict(span) -> dict:
    return {
        "story": span.story,
        "text": span.text,
        "anchor": span.anchor.to_dict(),
        "in_insert": span.in_insert,
        "in_delete": span.in_delete,
        "in_content_control": span.in_content_control,
        "in_text_box": span.in_text_box,
        "in_field": span.in_field,
        "crosses_paragraphs": span.crosses_paragraphs,
        "match_policy": span.match_policy,
        "raw_start": span._raw_start,
        "match_start": span._match_start,
    }


def op_outline(request: dict) -> dict:
    import docx
    from docx.story import outline

    document = docx.Document(request["path"])
    return {
        "outcome": "ok",
        "outlines": {
            view: outline(document, view=view).to_dict() for view in request["views"]
        },
    }


def op_find_many(request: dict) -> dict:
    import docx
    from docx.search import find_text

    document = docx.Document(request["path"])
    results = []
    for query in request["queries"]:
        try:
            spans = find_text(
                document,
                query["needle"],
                view=query.get("view", "current"),
                match=query.get("match", "exact"),
                story=query.get("story"),
                nth=query.get("nth"),
                near=query.get("near"),
            )
            results.append({"outcome": "ok", "spans": [_span_dict(s) for s in spans]})
        except Exception as exc:  # noqa: BLE001
            results.append(_error(exc))
    return {"outcome": "ok", "results": results}


def op_replace_cases(request: dict) -> dict:
    import docx
    from docx.package import patch_save
    from docx.search import find_one, replace_all

    results = []
    for case in request["cases"]:
        try:
            document = docx.Document(request["path"])
            if case.get("acknowledge"):
                from docx.protection import acknowledge_protection

                acknowledge_protection(document)
            tracked_kwargs = {}
            if case.get("tracked"):
                import datetime as _dt

                tracked_kwargs = {
                    "tracked": True,
                    "author": case.get("author"),
                    "date": _dt.datetime.fromisoformat(case["date"].replace("Z", "+00:00"))
                    if case.get("date")
                    else None,
                }
            if case.get("preserve_revision"):
                tracked_kwargs["preserve_revision"] = True
            if case.get("all"):
                result = replace_all(
                    document,
                    case["needle"],
                    case["with"],
                    match=case.get("match", "exact"),
                    view=case.get("view", "current"),
                    **tracked_kwargs,
                )
                payload = result.to_dict()
            else:
                span = find_one(
                    document,
                    case["needle"],
                    match=case.get("match", "exact"),
                    nth=case.get("nth"),
                    view=case.get("view", "current"),
                )
                payload = span.replace(case["with"], **tracked_kwargs).to_dict()
                if case.get("again"):
                    try:
                        span.replace(case["again"], **tracked_kwargs)
                        payload["again"] = "ok"
                    except Exception as exc:  # noqa: BLE001
                        payload["again"] = _category(exc)
            patch = patch_save(request["path"], document, case["out"]).to_dict()
            results.append({"outcome": "ok", "result": payload, "patch": patch})
        except Exception as exc:  # noqa: BLE001
            results.append(_error(exc))
    return {"outcome": "ok", "results": results}


def _extensions():
    """Per-work-package oracle extensions: tools/oracle/ext_*.py may define
    RESOLVERS ({"$tag": fn(document, value)}) and ENCODERS ([fn(value) ->
    (handled, encoded)]). Loaded by path because the oracle runs with -I."""
    import glob
    import importlib.util
    import os

    loaded = []
    here = os.path.dirname(os.path.abspath(__file__))
    for path in sorted(glob.glob(os.path.join(here, "ext_*.py"))):
        spec = importlib.util.spec_from_file_location(os.path.basename(path)[:-3], path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        loaded.append(module)
    return loaded


_EXTENSIONS = None


def _ext():
    global _EXTENSIONS
    if _EXTENSIONS is None:
        _EXTENSIONS = _extensions()
    return _EXTENSIONS


def _resolve_arg(document, value):
    """Decode `$`-tagged handle arguments used by shared ops cases."""
    import datetime as dt
    import importlib

    if isinstance(value, dict):
        for module in _ext():
            for tag, resolver in getattr(module, "RESOLVERS", {}).items():
                if tag in value:
                    return resolver(document, value)
    if isinstance(value, list):
        return [_resolve_arg(document, item) for item in value]
    if not isinstance(value, dict):
        return value
    if "$span" in value:
        from docx.search import find_one

        return find_one(
            document,
            value["$span"],
            nth=value.get("nth"),
            match=value.get("match", "exact"),
            story=value.get("story"),
            view=value.get("view", "current"),
        )
    if "$paragraph" in value:
        from docx.search import find_one
        from docx.text.paragraph import Paragraph

        span = find_one(document, value["$paragraph"], nth=value.get("nth"))
        element = span._atoms[0].paragraph
        return Paragraph(element, document._body)
    if "$block" in value:
        from docx.story import iter_blocks

        wanted = value["$block"]
        view = value.get("view", "current")
        matches = [b for b in iter_blocks(document, view=view) if b.text == wanted]
        index = value.get("index", 0)
        if index >= len(matches):
            raise LookupError(f"no block with text {wanted!r}")
        return matches[index]
    if "$table" in value:
        from docx.tableops import find_table

        return find_table(document, near_text=value["$table"], match=value.get("match", "exact"))
    if "$date" in value:
        return dt.datetime.fromisoformat(value["$date"].replace("Z", "+00:00"))
    if "$new" in value:
        module, _, name = value["$new"].rpartition(".")
        cls = getattr(importlib.import_module("docx." + module), name)
        kwargs = {k: _resolve_arg(document, v) for k, v in value.get("fields", {}).items()}
        return cls(**kwargs)
    return {k: _resolve_arg(document, v) for k, v in value.items()}


def _encode_result(value):
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    for module in _ext():
        for encoder in getattr(module, "ENCODERS", []):
            handled, encoded = encoder(value)
            if handled:
                return encoded
    if hasattr(value, "to_dict"):
        return value.to_dict()
    if isinstance(value, (list, tuple)):
        return [_encode_result(item) for item in value]
    if isinstance(value, dict):
        return {str(k): _encode_result(v) for k, v in value.items()}
    kind = type(value).__name__
    if kind == "Hyperlink":
        return {"type": "Hyperlink", "address": value.address, "text": value.text}
    return {"type": kind}


def op_ops(request: dict) -> dict:
    """Run one ops case: fresh document, steps in order, then patch_save."""
    import importlib
    import inspect

    import docx
    from docx.package import patch_save

    document = docx.Document(request["path"])
    steps = []
    for step in request["steps"]:
        try:
            module, _, name = step["op"].rpartition(".")
            if module == "ext":
                # Helper functions from tools/oracle/ext_*.py (FUNCTIONS dict);
                # they always receive the document first.
                target = next(
                    m.FUNCTIONS[name] for m in _ext() if name in getattr(m, "FUNCTIONS", {})
                )
                args = [document] + [_resolve_arg(document, a) for a in step.get("args", [])]
                kwargs = {k: _resolve_arg(document, v) for k, v in step.get("kwargs", {}).items()}
                steps.append({"outcome": "ok", "value": _encode_result(target(*args, **kwargs))})
                continue
            target = importlib.import_module("docx." + module)
            for attribute in name.split(":"):
                target = getattr(target, attribute)
            args = [_resolve_arg(document, a) for a in step.get("args", [])]
            kwargs = {k: _resolve_arg(document, v) for k, v in step.get("kwargs", {}).items()}
            parameters = list(inspect.signature(target).parameters)
            if parameters and parameters[0] == "document" and step.get("document", True):
                args.insert(0, document)
            steps.append({"outcome": "ok", "value": _encode_result(target(*args, **kwargs))})
        except Exception as exc:  # noqa: BLE001
            steps.append(_error(exc))
    patch = patch_save(request["path"], document, request["out"]).to_dict()
    return {"outcome": "ok", "steps": steps, "patch": patch}


def op_difflib(request: dict) -> dict:
    import difflib

    return {
        "outcome": "ok",
        "results": [
            [list(op) for op in difflib.SequenceMatcher(None, a, b).get_opcodes()]
            for a, b in request["pairs"]
        ],
    }


def op_text_diff(request: dict) -> dict:
    from docx.package import text_diff

    return {"outcome": "ok", "diff": text_diff(request["a"], request["b"], view=request.get("view", "current")).to_dict()}


def op_pending_changes(request: dict) -> dict:
    from docx.package import pending_changes

    return {"outcome": "ok", "diff": pending_changes(request["path"]).to_dict()}


def op_compare(request: dict) -> dict:
    import datetime as dt

    from docx.package import compare, patch_save

    result = compare(
        request["original"],
        request["revised"],
        author=request["author"],
        date=dt.datetime.fromisoformat(request["date"].replace("Z", "+00:00")),
        granularity=request.get("granularity", "word"),
        materialize=request.get("materialize"),
    )
    patch = patch_save(request["original"], result.document, request["out"]).to_dict()
    return {"outcome": "ok", "result": result.to_dict(), "patch": patch}


OPS = {
    "compare": op_compare,
    "difflib": op_difflib,
    "text_diff": op_text_diff,
    "pending_changes": op_pending_changes,
    "ops": op_ops,
    "replace_cases": op_replace_cases,
    "outline": op_outline,
    "find_many": op_find_many,
    "open": op_open,
    "diff": op_diff,
    "noop_patch_save": op_noop_patch_save,
    "diagnose": op_diagnose,
}


def main() -> None:
    with open(sys.argv[1], encoding="utf-8") as handle:
        requests = json.load(handle)
    results = []
    for request in requests:
        handler = OPS.get(request.get("op"))
        if handler is None:
            results.append({"outcome": "unsupported_op"})
            continue
        try:
            results.append(handler(request))
        except Exception as exc:  # noqa: BLE001 - every outcome is data here
            result = _error(exc)
            if request.get("trace"):
                result["trace"] = traceback.format_exc()
            results.append(result)
    json.dump(results, sys.stdout, indent=1, sort_keys=True)


if __name__ == "__main__":
    main()
