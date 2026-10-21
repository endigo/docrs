"""p7 oracle helpers: paragraph format, tab stops, fonts, inline shapes,
comment ranges and pictures from a path (Rust adapters:
ops/inherited/text.rs)."""

import p7_common as c


def p7_parfmt(document, para):
    return c.parfmt_values(c.paragraph(document, para).paragraph_format)


def p7_set_parfmt(document, para, props):
    fmt = c.paragraph(document, para).paragraph_format
    c.apply_parfmt(fmt, props)
    return c.parfmt_values(fmt)


def p7_font(document, para, run):
    return c.font_values(c.run(document, para, run).font)


def p7_set_font(document, para, run, changes):
    font = c.run(document, para, run).font
    c.apply_font(font, changes)
    return c.font_values(font)


def _tabs(tab_stops):
    return [
        {"position": int(t.position), "alignment": c.xml(t.alignment), "leader": c.xml(t.leader)}
        for t in tab_stops
    ]


def p7_tab_stops(document, para):
    """Reads `paragraph_format.tab_stops` (upstream adds `w:pPr` on access)."""
    return _tabs(c.paragraph(document, para).paragraph_format.tab_stops)


def p7_add_tab_stop(document, para, position, alignment="left", leader="none"):
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Emu

    tab_stops = c.paragraph(document, para).paragraph_format.tab_stops
    tab = tab_stops.add_tab_stop(
        Emu(position), WD_TAB_ALIGNMENT.from_xml(alignment), WD_TAB_LEADER.from_xml(leader)
    )
    return {"added": _tabs([tab])[0], "tabs": _tabs(tab_stops)}


def p7_clear_tab_stops(document, para):
    tab_stops = c.paragraph(document, para).paragraph_format.tab_stops
    tab_stops.clear_all()
    return _tabs(tab_stops)


def p7_delete_tab_stop(document, para, index):
    tab_stops = c.paragraph(document, para).paragraph_format.tab_stops

    def run():
        del tab_stops[index]
        return _tabs(tab_stops)

    return c.error_class(run)


def p7_set_tab_stop(document, para, index, prop, value):
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Emu

    tab_stops = c.paragraph(document, para).paragraph_format.tab_stops
    tab = tab_stops[index]
    if prop == "position":
        tab.position = Emu(value)
    elif prop == "alignment":
        tab.alignment = WD_TAB_ALIGNMENT.from_xml(value)
    elif prop == "leader":
        tab.leader = None if value is None else WD_TAB_LEADER.from_xml(value)
    else:
        raise LookupError("unknown tab stop property %r" % prop)
    return {"tab": _tabs([tab])[0], "tabs": _tabs(tab_stops)}


def _shape(shape):
    return {"width": int(shape.width), "height": int(shape.height), "type": c.name(shape.type)}


def p7_inline_shapes(document):
    return [_shape(shape) for shape in document.inline_shapes]


def p7_set_inline_shape(document, index, width=None, height=None):
    from docx.shared import Emu

    shape = document.inline_shapes[index]

    def run():
        if width is not None:
            shape.width = Emu(width)
        if height is not None:
            shape.height = Emu(height)
        return _shape(shape)

    return c.error_class(run)


def p7_mark_comment_range(document, para, run, last_para, last_run, comment_id):
    first = c.run(document, para, run)
    first.mark_comment_range(c.run(document, last_para, last_run), comment_id)


def p7_add_picture_path(document, path, width=None, height=None):
    shape = document.add_picture(path, width, height)
    return {"width": int(shape.width), "height": int(shape.height)}


def _run_item(item):
    from docx.drawing import Drawing

    if isinstance(item, str):
        return {"text": item}
    if isinstance(item, Drawing):
        return {"drawing": True}
    return {"page_break": True}


def _fragment(get):
    try:
        fragment = get()
    except ValueError as exc:
        return {"error": type(exc).__name__}
    if fragment is None:
        return None
    return {"text": fragment.text, "runs": [run.text for run in fragment.runs]}


def p7_paragraph_content(document, para):
    """`iter_inner_content()` of a paragraph and of each of its runs, its
    hyperlinks and the fragments around each rendered page break."""
    from docx.text.hyperlink import Hyperlink

    paragraph = c.paragraph(document, para)
    items = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            items.append(
                {
                    "hyperlink": {
                        "address": item.address,
                        "fragment": item.fragment,
                        "url": item.url,
                        "text": item.text,
                        "runs": [run.text for run in item.runs],
                        "contains_page_break": item.contains_page_break,
                    }
                }
            )
        else:
            items.append(
                {
                    "run": {
                        "text": item.text,
                        "contains_page_break": item.contains_page_break,
                        "content": [_run_item(x) for x in item.iter_inner_content()],
                    }
                }
            )
    breaks = [
        {
            "preceding": _fragment(lambda: b.preceding_paragraph_fragment),
            "following": _fragment(lambda: b.following_paragraph_fragment),
        }
        for b in paragraph.rendered_page_breaks
    ]
    return {
        "contains_page_break": paragraph.contains_page_break,
        "hyperlinks": len(paragraph.hyperlinks),
        "items": items,
        "page_breaks": breaks,
    }


def p7_style_names(document, para):
    """`paragraph.style.name` and `run.style.name` of each run."""
    paragraph = c.paragraph(document, para)

    def name(style):
        return None if style is None else style.name

    return {"paragraph": name(paragraph.style), "runs": [name(run.style) for run in paragraph.runs]}


def p7_table_style_name(document, table):
    style = c.table(document, table).style
    return None if style is None else style.name


def _bits(value):
    """IEEE-754 bit pattern of a float (exact, unlike JSON float text)."""
    import struct

    return struct.unpack("<q", struct.pack("<d", value))[0]


def p7_length(document, unit, value):
    """`docx.shared.<unit>(value)` and its unit accessors (floats as bits)."""
    from docx import shared

    length = getattr(shared, unit)(value)
    return {
        "emu": int(length),
        "cm": _bits(length.cm),
        "inches": _bits(length.inches),
        "mm": _bits(length.mm),
        "pt": _bits(length.pt),
        "twips": length.twips,
    }


def p7_rgb(document, text):
    from docx.shared import RGBColor

    return str(RGBColor.from_string(text))


def p7_clear_body(document):
    document._body.clear_content()
    return len(document.paragraphs) + len(document.tables)


def p7_paragraph_add_run(document, para, text=None, style=None):
    return {"text": c.paragraph(document, para).add_run(text, style).text}


def p7_run_add_picture_path(document, para, run, path, width=None, height=None):
    shape = c.run(document, para, run).add_picture(path, width, height)
    return {"width": int(shape.width), "height": int(shape.height)}


def p7_set_hyperlink_address(document, para, index, value):
    link = c.paragraph(document, para).hyperlinks[index]
    link.address = value
    return {"address": link.address, "fragment": link.fragment, "url": link.url}


def p7_image_from_file(document, path):
    from docx.image.image import Image

    image = Image.from_file(path)

    def native(get):
        try:
            return int(get())
        except Exception:  # noqa: BLE001
            return None

    return {
        "content_type": image.content_type,
        "px_width": image.px_width,
        "px_height": image.px_height,
        "horz_dpi": image.horz_dpi,
        "vert_dpi": image.vert_dpi,
        "ext": image.ext,
        "filename": image.filename,
        "width": native(lambda: image.width),
        "height": native(lambda: image.height),
        "sha1": image.sha1,
    }
