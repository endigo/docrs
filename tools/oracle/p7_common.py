"""Shared locators and encoders for the p7 oracle helpers (loaded by
`ext_p7.py`; mirrored by `crates/docrs-compat/src/ops/inherited.rs`).

Locators:

* container: ``"body"`` / ``None`` (the document body), ``{"section": i,
  "part": "header"|"footer"|"even_page_header"|...}`` (a header/footer of
  ``document.sections[i]``), or ``{"cell": [table, row, col]}`` (a table
  cell; ``table`` is a table locator);
* table: an int index into ``document.tables`` or ``{"in": container,
  "index": n}``;
* paragraph: an int index into ``document.paragraphs`` (negative allowed),
  the exact text of the first such paragraph, or ``{"in": container,
  "index": n}``;
* run: ``(paragraph, n)`` with ``n`` an index into ``paragraph.runs``.

Encoders: enum members by ``xml_value`` (``WD_LINE_SPACING`` and the
non-XML enums ``MSO_COLOR_TYPE`` / ``WD_TABLE_DIRECTION`` by name), lengths
as ints, and the error class name for non-refusal exceptions caught by a
helper (``{"error": "InvalidSpanError"}``).
"""

FONT_FLAGS = (
    "all_caps",
    "bold",
    "complex_script",
    "cs_bold",
    "cs_italic",
    "double_strike",
    "emboss",
    "hidden",
    "italic",
    "imprint",
    "math",
    "no_proof",
    "outline",
    "rtl",
    "shadow",
    "small_caps",
    "snap_to_grid",
    "spec_vanish",
    "strike",
    "web_hidden",
)


def xml(value):
    return None if value is None else value.xml_value


def name(value):
    return None if value is None else value.name


def length(value):
    return None if value is None else int(value)


def container(document, key):
    if key is None or key == "body":
        return document._body  # the body BlockItemContainer (`Document._body`)
    if "section" in key:
        return getattr(document.sections[key["section"]], key["part"])
    if "cell" in key:
        t, r, c = key["cell"]
        return table(document, t).cell(r, c)
    raise LookupError("bad container locator %r" % (key,))


def table(document, key):
    if isinstance(key, int):
        return document.tables[key]
    return container(document, key["in"]).tables[key["index"]]


def paragraph(document, key):
    if isinstance(key, dict):
        return container(document, key["in"]).paragraphs[key["index"]]
    paragraphs = document.paragraphs
    if isinstance(key, int):
        return paragraphs[key]
    for candidate in paragraphs:
        if candidate.text == key:
            return candidate
    raise LookupError("no paragraph with text %r" % key)


def run(document, para, index):
    return paragraph(document, para).runs[index]


def font_values(font):
    underline = font.underline
    if underline is not None and not isinstance(underline, bool):
        underline = underline.xml_value
    color = font.color
    rgb = color.rgb
    return {
        "flags": {flag: getattr(font, flag) for flag in FONT_FLAGS},
        "name": font.name,
        "size": length(font.size),
        "underline": underline,
        "color": {
            "type": name(color.type),
            "rgb": None if rgb is None else str(rgb),
            "theme_color": xml(color.theme_color),
        },
        "highlight": xml(font.highlight_color),
        "subscript": font.subscript,
        "superscript": font.superscript,
    }


def parfmt_values(fmt):
    line_spacing = fmt.line_spacing
    if line_spacing is not None:
        line_spacing = float(line_spacing) if isinstance(line_spacing, float) else int(line_spacing)
    return {
        "alignment": xml(fmt.alignment),
        "left_indent": length(fmt.left_indent),
        "right_indent": length(fmt.right_indent),
        "first_line_indent": length(fmt.first_line_indent),
        "space_before": length(fmt.space_before),
        "space_after": length(fmt.space_after),
        "line_spacing": line_spacing,
        "line_spacing_rule": name(fmt.line_spacing_rule),
        "keep_together": fmt.keep_together,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
        "widow_control": fmt.widow_control,
    }


_FORMAT_ORDER = (
    "alignment",
    "left_indent",
    "right_indent",
    "first_line_indent",
    "space_before",
    "space_after",
    "line_spacing",
    "line_spacing_rule",
    "keep_together",
    "keep_with_next",
    "page_break_before",
    "widow_control",
)


def apply_parfmt(fmt, props):
    """Assign `props` ({name: json value}) to a ParagraphFormat in the
    canonical order (the Rust `ParagraphFormat` field order)."""
    from docx.enum.text import WD_LINE_SPACING, WD_PARAGRAPH_ALIGNMENT
    from docx.shared import Emu

    unknown = set(props) - set(_FORMAT_ORDER)
    if unknown:
        raise LookupError("unknown paragraph format properties %r" % sorted(unknown))
    for prop in _FORMAT_ORDER:
        if prop not in props:
            continue
        value = props[prop]
        if value is not None:
            if prop == "alignment":
                value = WD_PARAGRAPH_ALIGNMENT.from_xml(value)
            elif prop == "line_spacing":
                value = float(value["multiple"]) if "multiple" in value else Emu(value["exact"])
            elif prop == "line_spacing_rule":
                value = WD_LINE_SPACING[value]
            elif prop in ("left_indent", "right_indent", "first_line_indent", "space_before", "space_after"):
                value = Emu(value)
        setattr(fmt, prop, value)


def apply_font(font, changes):
    """Apply `[name, value]` font assignments in order."""
    from docx.enum.dml import MSO_THEME_COLOR
    from docx.enum.text import WD_COLOR_INDEX, WD_UNDERLINE
    from docx.shared import Emu, RGBColor

    for prop, value in changes:
        if prop == "size":
            font.size = None if value is None else Emu(value)
        elif prop == "underline":
            font.underline = WD_UNDERLINE.from_xml(value) if isinstance(value, str) else value
        elif prop == "color":
            font.color.rgb = None if value is None else RGBColor.from_string(value)
        elif prop == "theme_color":
            font.color.theme_color = None if value is None else MSO_THEME_COLOR.from_xml(value)
        elif prop == "highlight":
            font.highlight_color = None if value is None else WD_COLOR_INDEX.from_xml(value)
        else:
            if not hasattr(type(font), prop):
                raise LookupError("unknown font property %r" % prop)
            setattr(font, prop, value)


def error_class(fn):
    """Run `fn`; a non-refusal, non-TypeError/ValueError exception is
    reported as {"error": class name} (the Rust adapters map their typed
    reasons to the same class names)."""
    from docx.errors import PaperRefusal

    try:
        return fn()
    except (PaperRefusal, TypeError, ValueError):
        raise
    except Exception as exc:  # noqa: BLE001
        return {"error": type(exc).__name__}
