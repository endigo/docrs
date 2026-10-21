"""p7 oracle helpers: sections, headers/footers, document settings and
block containers (Rust adapters: ops/inherited/sections.rs)."""

import io
import os

import p7_common as c

_HERE = os.path.dirname(os.path.abspath(__file__))
_TEST_FILES = os.path.join(
    _HERE, "..", "..", ".reference", "paper-docx", "features", "steps", "test_files"
)


def _test_file(document, value):
    return os.path.abspath(os.path.join(_TEST_FILES, value["$p7_test_file"]))


RESOLVERS = {"$p7_test_file": _test_file}

_LENGTHS = (
    "page_width",
    "page_height",
    "left_margin",
    "right_margin",
    "top_margin",
    "bottom_margin",
    "gutter",
    "header_distance",
    "footer_distance",
)


def _setup(section):
    out = {"start_type": c.xml(section.start_type), "orientation": c.xml(section.orientation)}
    for prop in _LENGTHS:
        out[prop] = c.length(getattr(section, prop))
    out["different_first_page_header_footer"] = section.different_first_page_header_footer
    return out


def _block_summary(item):
    from docx.table import Table

    if isinstance(item, Table):
        return {"kind": "tbl", "rows": len(item.rows), "cols": len(item.columns)}
    return {"kind": "p", "text": item.text, "style_id": item._p.style}


def p7_sections(document):
    """`len(document.sections)` and every section's page setup."""
    return [_setup(section) for section in document.sections]


def p7_section_set(document, index, changes):
    """Assign `[[prop, value], ...]` to `document.sections[index]` in order."""
    from docx.enum.section import WD_ORIENTATION, WD_SECTION_START
    from docx.shared import Emu

    section = document.sections[index]
    for prop, value in changes:
        if prop == "start_type":
            value = None if value is None else WD_SECTION_START.from_xml(value)
        elif prop == "orientation":
            value = None if value is None else WD_ORIENTATION.from_xml(value)
        elif prop in _LENGTHS:
            value = None if value is None else Emu(value)
        elif prop != "different_first_page_header_footer":
            raise LookupError("unknown section property %r" % prop)
        setattr(section, prop, value)
    return _setup(section)


def p7_section_content(document, index):
    return [_block_summary(item) for item in document.sections[index].iter_inner_content()]


def _hdrftr(document, index, part):
    return getattr(document.sections[index], part)


def p7_hdrftr_linked(document, index, part):
    return _hdrftr(document, index, part).is_linked_to_previous


def p7_hdrftr_set_linked(document, index, part, value):
    hdrftr = _hdrftr(document, index, part)
    hdrftr.is_linked_to_previous = value
    return hdrftr.is_linked_to_previous


def p7_container_content(document, container):
    """`iter_inner_content()` of a container (a header/footer adds its
    definition first, as upstream `_get_or_add_definition`)."""
    return [_block_summary(item) for item in c.container(document, container).iter_inner_content()]


def p7_container_add_paragraph(document, container, text="", style=None):
    return _block_summary(c.container(document, container).add_paragraph(text, style))


def p7_container_add_table(document, container, rows, cols, width=None):
    from docx.shared import Emu

    target = c.container(document, container)
    if width is None:
        table = target.add_table(rows, cols)  # _Cell.add_table(rows, cols)
    else:
        table = target.add_table(rows, cols, Emu(width))
    return _block_summary(table)


def p7_paragraph_set_text(document, para, text):
    paragraph = c.paragraph(document, para)
    paragraph.text = text
    return _block_summary(paragraph)


def p7_paragraph_set_style(document, para, style):
    paragraph = c.paragraph(document, para)
    paragraph.style = style
    return _block_summary(paragraph)


def p7_paragraph_add_picture(document, para, image, width=None, height=None):
    """`paragraph.add_run().add_picture(image)` (header/footer pictures)."""
    shape = c.paragraph(document, para).add_run().add_picture(io.BytesIO(image), width, height)
    return {"width": int(shape.width), "height": int(shape.height)}


def p7_odd_even(document):
    return document.settings.odd_and_even_pages_header_footer


def p7_set_odd_even(document, value):
    settings = document.settings
    settings.odd_and_even_pages_header_footer = value
    return settings.odd_and_even_pages_header_footer
