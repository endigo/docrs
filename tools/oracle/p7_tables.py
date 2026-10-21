"""p7 oracle helpers: table, row, column and cell properties and cell
merging (Rust adapters: ops/inherited/tables.rs)."""

import p7_common as c


def _direction(table):
    value = table.table_direction
    return None if value is None else ("RTL" if value else "LTR")


def _rows(table):
    return {
        "rows": [
            {
                "height": c.length(row.height),
                "height_rule": c.xml(row.height_rule),
                "grid_cols_before": row.grid_cols_before,
                "grid_cols_after": row.grid_cols_after,
                "cells": [cell.text for cell in row.cells],
            }
            for row in table.rows
        ],
        "columns": [
            {"width": c.length(col.width), "cells": [cell.text for cell in col.cells]}
            for col in table.columns
        ],
    }


def p7_rows(document, table):
    """Rows and columns only (no `w:tblPr` access)."""
    return _rows(c.table(document, table))


def _table(table):
    return {
        "alignment": c.xml(table.alignment),
        "autofit": table.autofit,
        "direction": _direction(table),
        "style_id": table._tbl.tblStyle_val,
        "rows": [
            {
                "height": c.length(row.height),
                "height_rule": c.xml(row.height_rule),
                "grid_cols_before": row.grid_cols_before,
                "grid_cols_after": row.grid_cols_after,
                "cells": [cell.text for cell in row.cells],
            }
            for row in table.rows
        ],
        "columns": [
            {"width": c.length(col.width), "cells": [cell.text for cell in col.cells]}
            for col in table.columns
        ],
    }


def _cell(cell):
    return {
        "text": cell.text,
        "width": c.length(cell.width),
        "grid_span": cell.grid_span,
        "vertical_alignment": c.xml(cell.vertical_alignment),
        "paragraphs": len(cell.paragraphs),
        "tables": len(cell.tables),
    }


def p7_table(document, table):
    return _table(c.table(document, table))


def p7_table_set(document, table, changes):
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_TABLE_DIRECTION

    target = c.table(document, table)
    for prop, value in changes:
        if prop == "alignment":
            target.alignment = None if value is None else WD_TABLE_ALIGNMENT.from_xml(value)
        elif prop == "autofit":
            target.autofit = value
        elif prop == "direction":
            target.table_direction = None if value is None else WD_TABLE_DIRECTION[value]
        else:
            raise LookupError("unknown table property %r" % prop)
    return _table(target)


def p7_row_set(document, table, row, changes):
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.shared import Emu

    target = c.table(document, table)
    tr = target.rows[row]
    for prop, value in changes:
        if prop == "height":
            tr.height = None if value is None else Emu(value)
        elif prop == "height_rule":
            tr.height_rule = None if value is None else WD_ROW_HEIGHT_RULE.from_xml(value)
        else:
            raise LookupError("unknown row property %r" % prop)
    return _rows(target)


def p7_column_set_width(document, table, col, width):
    from docx.shared import Emu

    target = c.table(document, table)
    target.columns[col].width = None if width is None else Emu(width)
    return _rows(target)


def p7_cells(document, table):
    """Every layout-grid cell, row by row (`row_cells(i)` for each row)."""
    target = c.table(document, table)
    return [[_cell(cell) for cell in target.row_cells(i)] for i in range(len(target.rows))]


def p7_column_cells(document, table, col):
    return [cell.text for cell in c.table(document, table).column_cells(col)]


def p7_cell_set(document, table, row, col, changes):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Emu

    cell = c.table(document, table).cell(row, col)
    for prop, value in changes:
        if prop == "width":
            cell.width = Emu(value)
        elif prop == "vertical_alignment":
            cell.vertical_alignment = (
                None if value is None else WD_CELL_VERTICAL_ALIGNMENT.from_xml(value)
            )
        else:
            raise LookupError("unknown cell property %r" % prop)
    return _cell(cell)


def p7_merge(document, table, row, col, row2, col2):
    target = c.table(document, table)

    def run():
        merged = target.cell(row, col).merge(target.cell(row2, col2))
        return {"merged": _cell(merged), "table": _table(target)}

    return c.error_class(run)
