"""Oracle extension for M4 work package p4 (inherited authoring API, images,
picture replacement).

RESOLVERS: {"$tag": fn(document, value)} argument decoders.
ENCODERS: [fn(value) -> (handled, encoded)] result encoders.
FUNCTIONS: {"name": fn(document, *args, **kwargs)} callable as op "ext.name".

Locators shared with the Rust adapters (crates/docrs-compat/src/ops/
authoring.rs, drawing.rs):

* paragraph: an int index into `document.paragraphs` (negative allowed) or
  the exact text of the first such paragraph;
* run: an int index into `paragraph.runs`;
* table: an int index into `document.tables`;
* drawing: an int index into the main part's `w:drawing` elements in
  document order, or {"part": "word/header1.xml", "index": n};
* image bytes: {"$p4_image": "<hex>"} or {"$p4_image_file": "<name>"} (a file
  in the pinned checkout's features/steps/test_files).
"""

import io
import os
import types

_HERE = os.path.dirname(os.path.abspath(__file__))
_TEST_FILES = os.path.join(
    _HERE, "..", "..", ".reference", "paper-docx", "features", "steps", "test_files"
)


def _image_hex(document, value):
    return bytes.fromhex(value["$p4_image"])


def _image_file(document, value):
    with open(os.path.join(_TEST_FILES, value["$p4_image_file"]), "rb") as handle:
        return handle.read()


RESOLVERS = {"$p4_image": _image_hex, "$p4_image_file": _image_file}
ENCODERS = []


def _para(document, key):
    paragraphs = document.paragraphs
    if isinstance(key, int):
        return paragraphs[key]
    for paragraph in paragraphs:
        if paragraph.text == key:
            return paragraph
    raise LookupError("no paragraph with text %r" % key)


def _run(document, para, run):
    return _para(document, para).runs[run]


def _table(document, key):
    return document.tables[key]


def _para_summary(paragraph):
    return {"text": paragraph.text, "style_id": paragraph._p.style}


def _table_summary(table):
    tbl = table._tbl
    return {
        "rows": len(tbl.tr_lst),
        "cols": len(tbl.tblGrid.gridCol_lst),
        "style_id": tbl.tblStyle_val,
    }


def _try_int(fn):
    try:
        return int(fn())
    except Exception:  # noqa: BLE001
        return None


def _image_summary(image):
    return {
        "content_type": image.content_type,
        "px_width": image.px_width,
        "px_height": image.px_height,
        "horz_dpi": image.horz_dpi,
        "vert_dpi": image.vert_dpi,
        "ext": image.ext,
        "filename": image.filename,
        "width": _try_int(lambda: image.width),
        "height": _try_int(lambda: image.height),
        "sha1": image.sha1,
    }


def p4_characterize(document, image):
    from docx.image.image import Image

    try:
        return _image_summary(Image.from_blob(image))
    except Exception as exc:  # noqa: BLE001
        return {"error": type(exc).__name__}


def p4_scaled(document, image, width=None, height=None):
    from docx.image.image import Image

    try:
        cx, cy = Image.from_blob(image).scaled_dimensions(width, height)
        return [int(cx), int(cy)]
    except Exception as exc:  # noqa: BLE001
        return {"error": type(exc).__name__}


def p4_add_paragraph(document, text="", style=None):
    return _para_summary(document.add_paragraph(text, style))


def p4_add_heading(document, text="", level=1):
    return _para_summary(document.add_heading(text, level))


def p4_add_page_break(document):
    return _para_summary(document.add_page_break())


def p4_add_table(document, rows, cols, style=None):
    return _table_summary(document.add_table(rows, cols, style))


def _shape(shape):
    return {"width": int(shape.width), "height": int(shape.height)}


def p4_add_picture(document, image, width=None, height=None):
    return _shape(document.add_picture(io.BytesIO(image), width, height))


def p4_add_section(document, start_type="nextPage"):
    from docx.enum.section import WD_SECTION_START

    value = WD_SECTION_START.from_xml(start_type)
    return {"start_type": document.add_section(value).start_type.xml_value}


def p4_add_run(document, para, text=None, style=None):
    return {"text": _para(document, para).add_run(text, style).text}


def p4_insert_paragraph_before(document, para, text=None, style=None):
    return _para_summary(_para(document, para).insert_paragraph_before(text, style))


def p4_set_paragraph_text(document, para, text):
    paragraph = _para(document, para)
    paragraph.text = text
    return _para_summary(paragraph)


def p4_clear_paragraph(document, para):
    return _para_summary(_para(document, para).clear())


def p4_set_paragraph_style(document, para, style):
    paragraph = _para(document, para)
    paragraph.style = style
    return _para_summary(paragraph)


def _alignment(value):
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    return None if value is None else WD_PARAGRAPH_ALIGNMENT.from_xml(value)


def p4_set_alignment(document, para, value):
    _para(document, para).alignment = _alignment(value)


_FORMAT_ORDER = (
    "alignment",
    "left_indent",
    "right_indent",
    "first_line_indent",
    "space_before",
    "space_after",
    "line_spacing",
    "keep_together",
    "keep_with_next",
    "page_break_before",
    "widow_control",
)


def p4_paragraph_format(document, para, props):
    from docx.shared import Emu

    unknown = set(props) - set(_FORMAT_ORDER)
    if unknown:
        raise LookupError("unknown paragraph format properties %r" % sorted(unknown))
    fmt = _para(document, para).paragraph_format
    for name in _FORMAT_ORDER:
        if name not in props:
            continue
        value = props[name]
        if name == "alignment":
            value = _alignment(value)
        elif name == "line_spacing" and value is not None:
            value = float(value["multiple"]) if "multiple" in value else Emu(value["exact"])
        elif name in ("left_indent", "right_indent", "first_line_indent", "space_before", "space_after") and value is not None:
            value = Emu(value)
        setattr(fmt, name, value)


def p4_set_run_text(document, para, run, text):
    target = _run(document, para, run)
    target.text = text
    return {"text": target.text}


def p4_run_add_text(document, para, run, text):
    target = _run(document, para, run)
    target.add_text(text)
    return {"text": target.text}


def p4_run_add_tab(document, para, run):
    target = _run(document, para, run)
    target.add_tab()
    return {"text": target.text}


def p4_run_add_break(document, para, run, kind="line"):
    from docx.enum.text import WD_BREAK

    kinds = {
        "line": WD_BREAK.LINE,
        "page": WD_BREAK.PAGE,
        "column": WD_BREAK.COLUMN,
        "line_clear_left": WD_BREAK.LINE_CLEAR_LEFT,
        "line_clear_right": WD_BREAK.LINE_CLEAR_RIGHT,
        "line_clear_all": WD_BREAK.LINE_CLEAR_ALL,
    }
    target = _run(document, para, run)
    target.add_break(kinds[kind])
    return {"text": target.text}


def p4_run_clear(document, para, run):
    return {"text": _run(document, para, run).clear().text}


def p4_set_run_style(document, para, run, style):
    _run(document, para, run).style = style


def p4_run_font(document, para, run, changes):
    """Apply `[name, value]` font assignments in order."""
    from docx.enum.text import WD_COLOR_INDEX, WD_UNDERLINE
    from docx.shared import Emu, RGBColor

    font = _run(document, para, run).font
    for name, value in changes:
        if name == "size":
            font.size = None if value is None else Emu(value)
        elif name == "underline":
            font.underline = WD_UNDERLINE.from_xml(value) if isinstance(value, str) else value
        elif name == "color":
            font.color.rgb = None if value is None else RGBColor.from_string(value)
        elif name == "highlight":
            font.highlight_color = None if value is None else WD_COLOR_INDEX.from_xml(value)
        else:
            if not hasattr(type(font), name):
                raise LookupError("unknown font property %r" % name)
            setattr(font, name, value)


def p4_run_add_picture(document, para, run, image, width=None, height=None):
    return _shape(_run(document, para, run).add_picture(io.BytesIO(image), width, height))


def p4_cell_set_text(document, table, row, col, text):
    cell = _table(document, table).cell(row, col)
    cell.text = text
    return {"text": cell.text}


def p4_table_add_row(document, table):
    target = _table(document, table)
    target.add_row()
    return _table_summary(target)


def p4_table_add_column(document, table, width):
    from docx.shared import Emu

    target = _table(document, table)
    target.add_column(Emu(width))
    return _table_summary(target)


def p4_table_set_style(document, table, style):
    target = _table(document, table)
    target.style = style
    return _table_summary(target)


def _drawing(document, key):
    from docx.drawing import Drawing
    from docx.oxml.ns import qn

    if isinstance(key, int):
        part, index = document.part, key
    else:
        wanted = "/" + key["part"].lstrip("/")
        part = next(p for p in document.part.package.iter_parts() if p.partname == wanted)
        index = key.get("index", 0)
    element = list(part.element.iter(qn("w:drawing")))[index]
    return Drawing(element, types.SimpleNamespace(part=part))


def p4_has_picture(document, drawing):
    return _drawing(document, drawing).has_picture


def p4_drawing_image(document, drawing):
    target = _drawing(document, drawing)
    try:
        image = target.image
    except Exception as exc:  # noqa: BLE001
        return {"error": type(exc).__name__}
    return _image_summary(image)


def p4_replace_picture(document, drawing, image):
    _drawing(document, drawing).replace_picture(io.BytesIO(image))


FUNCTIONS = {
    name: fn
    for name, fn in dict(globals()).items()
    if name.startswith("p4_") and callable(fn)
}
