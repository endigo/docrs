"""Reopen exported .docx files with independent readers.

Usage: python -I tools/markdown/check_exports.py DIR [--word] [--png]

For every DIR/*.docx:
  * python-docx (the pinned upstream oracle under .reference/) opens it and
    lists paragraph styles, tables and inline shapes;
  * pandoc (when installed) reads it back to GitHub Markdown;
  * with --word, Microsoft Word (macOS) opens it and exports a PDF, which
    proves Word loads the file without a repair prompt; with --png the PDF
    pages are rendered for viewing (needs swiftc).
Writes DIR/report.json and prints one line per file. Exit status 1 when any
reader fails.
"""

import json
import os
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, ".reference", "paper-docx", "src"))
import docx  # noqa: E402  (the pinned oracle)

HERE = os.path.dirname(os.path.abspath(__file__))
WORD_TMP = os.path.expanduser("~/Library/Containers/com.microsoft.Word/Data/tmp/docrs-check")


def python_docx(path):
    document = docx.Document(path)
    return {
        "paragraphs": [[p.style.name, p.text] for p in document.paragraphs],
        "tables": [[[c.text for c in row.cells] for row in t.rows] for t in document.tables],
        "inline_shapes": len(document.inline_shapes),
        "sections": len(document.sections),
    }


def pandoc(path):
    if not shutil.which("pandoc"):
        return None
    result = subprocess.run(
        ["pandoc", "-f", "docx", "-t", "gfm", "--wrap=none", path],
        capture_output=True, text=True, timeout=120,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip())
    return result.stdout


def word_pdf(path, out_dir, png):
    os.makedirs(WORD_TMP, exist_ok=True)
    base = os.path.splitext(os.path.basename(path))[0]
    staged = os.path.join(WORD_TMP, base + ".docx")
    pdf = os.path.join(WORD_TMP, base + ".pdf")
    shutil.copy(path, staged)
    if os.path.exists(pdf):
        os.remove(pdf)
    subprocess.run(
        ["osascript", os.path.join(HERE, "word2pdf.applescript"), staged, pdf],
        check=True, capture_output=True, text=True, timeout=180,
    )
    target = os.path.join(out_dir, base + ".pdf")
    shutil.copy(pdf, target)
    pages = None
    if png:
        tool = os.path.join(out_dir, ".pdfpages")
        if not os.path.exists(tool):
            subprocess.run(["swiftc", "-O", os.path.join(HERE, "pdfpages.swift"), "-o", tool],
                           check=True, capture_output=True)
        result = subprocess.run([tool, target, os.path.join(out_dir, base), "1.0"],
                                check=True, capture_output=True, text=True)
        pages = int(result.stdout.split(":")[1])
    return {"pdf": target, "pages": pages}


def main():
    directory = sys.argv[1]
    use_word = "--word" in sys.argv
    png = "--png" in sys.argv
    report = {}
    failed = False
    for name in sorted(os.listdir(directory)):
        if not name.endswith(".docx"):
            continue
        path = os.path.join(directory, name)
        entry = {}
        for reader, run in (("python_docx", lambda: python_docx(path)),
                            ("pandoc", lambda: pandoc(path)),
                            ("word", (lambda: word_pdf(path, directory, png)) if use_word else (lambda: None))):
            try:
                entry[reader] = run()
            except Exception as error:  # report every reader failure
                entry[reader] = {"error": str(error)}
                failed = True
        report[name] = entry
        status = ", ".join(
            f"{reader}: {'FAIL ' + value['error'][:80] if isinstance(value, dict) and 'error' in value else ('skipped' if value is None else 'ok')}"
            for reader, value in entry.items()
        )
        pages = entry.get("word") or {}
        if isinstance(pages, dict) and pages.get("pages"):
            status += f" ({pages['pages']} pages)"
        print(f"{name}: {status}")
    with open(os.path.join(directory, "report.json"), "w") as handle:
        json.dump(report, handle, indent=1, ensure_ascii=False)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
